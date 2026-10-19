#!/usr/bin/env python3
"""
Deterministic Institutional Word (.docx) Memorandum Generator
(`scripts/generate_docx_memo.py`).

Generates `deliverables/Инвестиционный_меморандум_MOEX_SPBEX_2026.docx` from the
verified quantitative matrices (`data/processed/*.json`), pairing Georgia editorial
serif headings with Arial body/tabular text, institutional slate headers (#F4F5F7),
and pure Russian financial typography.
"""
from __future__ import annotations

import json
import pathlib
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

ROOT_DIR = pathlib.Path(__file__).resolve().parent.parent


def fmt_ru(val: float | int | None, decimals: int = 1) -> str:
    if val is None:
        return "—"
    s = f"{float(val):,.{decimals}f}"
    parts = s.split(".")
    parts[0] = parts[0].replace(",", "\u00a0")
    return ",".join(parts) if decimals > 0 else parts[0]


def set_cell_shading(cell, color_hex: str) -> None:
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_margins(cell, top: int = 90, bottom: int = 90, left: int = 120, right: int = 120) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for side, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def add_heading_georgia(doc: Document, text: str, size_pt: int = 15, space_before: int = 16, space_after: int = 6) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Georgia"
    run.font.size = Pt(size_pt)
    run.font.bold = True
    run.font.color.rgb = RGBColor(22, 50, 91)


def add_body_para(doc: Document, text: str, bold_prefix: str = "") -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.18
    if bold_prefix:
        r0 = p.add_run(bold_prefix)
        r0.font.name = "Arial"
        r0.font.size = Pt(10.5)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(15, 23, 42)
    r1 = p.add_run(text)
    r1.font.name = "Arial"
    r1.font.size = Pt(10.5)
    r1.font.color.rgb = RGBColor(20, 23, 28)


def style_table(table, headers: list[str], rows: list[list[str]], right_align_cols: set[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))
    trPr.append(OxmlElement("w:cantSplit"))

    for idx, h_text in enumerate(headers):
        c = hdr_cells[idx]
        set_cell_shading(c, "F4F5F7")
        set_cell_margins(c, 100, 100, 110, 110)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if idx in right_align_cols else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(h_text)
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(22, 50, 91)

    for r_data in rows:
        row_obj = table.add_row()
        r_trPr = row_obj._tr.get_or_add_trPr()
        r_trPr.append(OxmlElement("w:cantSplit"))
        for idx, val_str in enumerate(r_data):
            c = row_obj.cells[idx]
            set_cell_margins(c, 75, 75, 110, 110)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if idx in right_align_cols else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(val_str)
            run.font.name = "Arial"
            run.font.size = Pt(8.5)
            if idx == 0 or idx == len(r_data) - 1:
                run.font.bold = True
            run.font.color.rgb = RGBColor(20, 23, 28)


def main() -> None:
    val_data = json.loads((ROOT_DIR / "data" / "processed" / "quant_valuation_matrix.json").read_text(encoding="utf-8"))
    sec_data = json.loads((ROOT_DIR / "data" / "processed" / "sector_aggregates.json").read_text(encoding="utf-8"))
    port_data = json.loads((ROOT_DIR / "data" / "processed" / "optimal_portfolios.json").read_text(encoding="utf-8"))
    exch_data = json.loads((ROOT_DIR / "data" / "raw" / "exchange_merged_snapshot.json").read_text(encoding="utf-8"))

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Institutional Masthead
    p_mast = doc.add_paragraph()
    p_mast.paragraph_format.space_after = Pt(4)
    r_mast = p_mast.add_run("АНАЛИТИЧЕСКИЙ МЕМОРАНДУМ ИНВЕСТИЦИОННОГО КОМИТЕТА")
    r_mast.font.name = "Georgia"
    r_mast.font.size = Pt(11)
    r_mast.font.bold = True
    r_mast.font.color.rgb = RGBColor(22, 50, 91)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(10)
    r_title = p_title.add_run(
        "Фундаментальный аудит отчётностей МСФО и РСБУ, отраслевой скоринг и аллокация капитала на Московской и СПБ биржах"
    )
    r_title.font.name = "Georgia"
    r_title.font.size = Pt(19)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    mp = exch_data["macro_market_pulse"]
    add_body_para(
        doc,
        f"Дата расчётного слепка: 19 октября 2026 г.  |  Ключевая ставка ЦБ РФ: 14,00% (с 27.07.2026)  |  "
        f"Курс ЦБ USD/RUB: 82,85 руб.  |  Индекс Мосбиржи (IMOEX): {fmt_ru(mp.get('imoex'), 2)} п.  |  "
        f"Индекс RGBI: {fmt_ru(mp.get('rgbi'), 2)} п.  |  Оборот MOEX TQBR: {fmt_ru(mp.get('moex_tqbr_turnover_bln_rub'), 2)} млрд руб.  |  "
        f"Оборот СПБ Биржи: {fmt_ru(mp.get('spbex_total_turnover_bln_rub'), 3)} млрд руб.",
    )

    add_heading_georgia(doc, "1. Траектория ключевой ставки ЦБ РФ (21,00% → 14,00%) и поляризация балансов")
    add_body_para(
        doc,
        "По данным официальной базы Банка России (cbr.ru/hd_base/KeyRate/), после пика 21,00% годовых в октябре 2024 года регулятор провёл десять последовательных снижений ставки суммарным шагом 700 б.п., установив с 27 июля 2026 года уровень 14,00% годовых (при доходности длинных ОФЗ около 13,4% и фондов ликвидности около 13,85% годовых). Вместе с тем эффективная стоимость корпоративного долга под плавающую ставку (КС + 2,2–3,5%) остаётся в диапазоне 16,2–17,5% годовых, продолжая поглощать основную часть операционной прибыли закредитованных компаний, тогда как эмитенты с отрицательным чистым долгом одновременно получают около 11,9% годовых по депозитам и выигрывают от расширения мультипликаторов.",
    )
    add_body_para(
        doc,
        "Детерминированная сверка 62 бухгалтерских тождеств МСФО и РСБУ по 31 ведущему эмитенту восьми отраслей и автоматический парсинг 120 живых новостей (ленты ЦБ РФ, Мосбиржи ISS и деловых СМИ) позволили объединить фундаментальную оценку с корпоративными катализаторами и разделить рекомендации по срокам инвестирования на три горизонта: тактический (1–3 месяца), базовый годовой (6–12 месяцев) и стратегический (1–3 года).",
    )

    add_heading_georgia(doc, "2. Отраслевая матрица фундаментальных мультипликаторов и Z-оценок (8 секторов)")
    sec_headers = [
        "Отрасль",
        "Бумаг",
        "Капитализ., млрд руб.",
        "EBITDA, млрд руб.",
        "Чист. прибыль, млрд руб.",
        "FCF, млрд руб.",
        "Чист. долг, млрд руб.",
        "Мед. P/E",
        "Сред. ДД",
        "Балл",
        "Фаворит",
    ]
    sec_rows = []
    for s in sec_data["sectors"]:
        sec_rows.append([
            s["sector"],
            str(s["companies_count"]),
            fmt_ru(s["total_mcap_bln_rub"], 1),
            fmt_ru(s["total_ebitda_bln_rub"], 1),
            fmt_ru(s["total_net_income_bln_rub"], 1),
            fmt_ru(s["total_fcf_bln_rub"], 1),
            fmt_ru(s["total_net_debt_bln_rub"], 1),
            f"{fmt_ru(s['median_pe'], 2)}×" if s["median_pe"] else "—",
            f"{fmt_ru(s['avg_div_yield_pct'], 1)}%",
            fmt_ru(s["avg_composite_score"], 1),
            f"{s['top_pick_ticker']} (+{fmt_ru(s['top_pick_return_pct'], 1)}%)",
        ])
    t_sec = doc.add_table(rows=1, cols=len(sec_headers))
    style_table(t_sec, sec_headers, sec_rows, {1, 2, 3, 4, 5, 6, 7, 8, 9})

    add_heading_georgia(doc, "3. Инвестиционный консультант по 3 горизонтам (1–3М / 6–12М / 1–3Г): ведущие бумаги")
    add_body_para(
        doc,
        "На горизонте 1–3 месяца приоритет отдаётся бумагам с зимними промежуточными дивидендными отсечками (ЛУКОЙЛ, ИКС 5, Хэдхантер, Татнефть) и лимитному выкупу эмоциональных октябрьских откатов (Полюс в коридоре 1 081,4–1 139,3 руб., Яндекс в коридоре 3 231,3–3 385,9 руб.). На горизонте 6–12 месяцев максимальную полную доходность генерируют денежные кубышки и банки с высоким запасом капитала (Сургутнефтегаз-п, Банк Санкт-Петербург, Транснефть-п, Мосбиржа, Интер РАО), а на горизонте 1–3 года — компаундеры с высокой рентабельностью капитала (Хэдхантер, ИКС 5, Т-Технологии, Яндекс, Полюс, Элемент).",
    )
    issuers_list = list(val_data["issuers"].values())
    top_subset = issuers_list[:11] + [val_data["issuers"]["T"], val_data["issuers"]["YDEX"], val_data["issuers"]["ELMT"]]
    iss_headers = [
        "Ранг",
        "Тикер",
        "Эмитент",
        "Цена, руб.",
        "Коридор входа",
        "Доходн. 1–3М",
        "Таргет 12М, руб.",
        "Доходн. 6–12М",
        "3Г CAGR / год",
        "P/E",
        "Балл",
        "Реком. срок",
    ]
    iss_rows = []
    for x in top_subset:
        dec = 4 if x["price_rub"] < 5 else 2
        h_s = x["horizons"]["short_1_3m"]
        h_m = x["horizons"]["medium_6_12m"]
        h_l = x["horizons"]["long_1_3y"]
        iss_rows.append([
            str(x["global_rank"]),
            x["ticker"],
            x["short_name"],
            fmt_ru(x["price_rub"], dec),
            f"{fmt_ru(h_s['entry_zone_rub'][0], dec)}–{fmt_ru(h_s['entry_zone_rub'][1], dec)}",
            f"+{fmt_ru(h_s['expected_total_return_pct'], 1)}%",
            fmt_ru(h_m["target_price_rub"], dec),
            f"+{fmt_ru(h_m['expected_total_return_pct'], 1)}%",
            f"+{fmt_ru(h_l['expected_cagr_3y_pct'], 1)}%",
            f"{fmt_ru(x['pe_ratio'], 2)}×" if x.get("pe_ratio") else "—",
            fmt_ru(x["composite_score"], 1),
            x["recommended_holding_months"],
        ])
    t_iss = doc.add_table(rows=1, cols=len(iss_headers))
    style_table(t_iss, iss_headers, iss_rows, {0, 3, 4, 5, 6, 7, 8, 9, 10})

    add_heading_georgia(doc, "4. Структура флагманского модельного портфеля на капитал 1 000 000 рублей")
    core_p = port_data["portfolios"][0]
    add_body_para(
        doc,
        f"Портфель «{core_p['name']}» (рекомендуемый срок удержания — {core_p['recommended_horizon']}) объединяет лидеров генерации денежного потока на Московской бирже, технологических эмитентов СПБ Биржи и ликвидный буфер денежного рынка. Ожидаемая полная доходность составляет +{fmt_ru(core_p['expected_1_3m_total_return_pct'], 2)}% на горизонте 1–3 месяца, +{fmt_ru(core_p['expected_12m_total_return_pct'], 2)}% на горизонте 6–12 месяцев (при дивидендном потоке {fmt_ru(core_p['expected_portfolio_div_yield_pct'], 2)}%, или {fmt_ru(core_p['expected_annual_div_rub_per_1m'], 0)} руб. в год на 1 млн руб.) и +{fmt_ru(core_p['expected_3y_cagr_pct'], 2)}% годовых на трёхлетнем горизонте при гармоническом P/E {fmt_ru(core_p['harmonic_pe_ratio'], 2)}×.",
    )
    p_headers = [
        "Тикер",
        "Инструмент",
        "Биржа",
        "Доля, %",
        "Сумма, руб.",
        "Цена, руб.",
        "Лотов",
        "Див. 12М",
        "Доходн. 1–3М",
        "Доходн. 6–12М",
        "3Г CAGR",
    ]
    p_rows = []
    for pos in core_p["positions"]:
        p_rows.append([
            pos["ticker"],
            pos["name"],
            pos["exchange"],
            f"{fmt_ru(pos['weight_pct'], 1)}%",
            fmt_ru(pos["target_allocation_rub"], 0),
            fmt_ru(pos["price_rub"], 4 if pos["price_rub"] < 5 else 2),
            fmt_ru(pos["recommended_lots_per_1m"], 0),
            f"{fmt_ru(pos['div_yield_pct'], 2)}%",
            f"+{fmt_ru(pos['expected_1_3m_return_pct'], 1)}%",
            f"+{fmt_ru(pos['expected_12m_return_pct'], 1)}%",
            f"+{fmt_ru(pos['expected_3y_cagr_pct'], 1)}%",
        ])
    t_port = doc.add_table(rows=1, cols=len(p_headers))
    style_table(t_port, p_headers, p_rows, {3, 4, 5, 6, 7, 8, 9, 10})

    out_docx = ROOT_DIR / "deliverables" / "Инвестиционный_меморандум_MOEX_SPBEX_2026.docx"
    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_docx))
    print(f"[OK] Saved institutional Word memorandum -> {out_docx}")


if __name__ == "__main__":
    main()
